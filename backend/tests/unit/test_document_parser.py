from __future__ import annotations

import io

import pytest
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, StreamObject

from app.core.exceptions import ExtractionError
from app.parsers.document_parser import parse_document, parse_docx, parse_pdf, parse_pptx, parse_txt
from shared.constants import ContentBlockType, ContentUnitType


def _build_pdf(text: str, pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    for page_index, page in enumerate(writer.pages):
        font = DictionaryObject(
            {
                NameObject("/F1"): DictionaryObject(
                    {
                        NameObject("/Type"): NameObject("/Font"),
                        NameObject("/Subtype"): NameObject("/Type1"),
                        NameObject("/BaseFont"): NameObject("/Helvetica"),
                    }
                )
            }
        )
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): font})
        stream = StreamObject()
        text = text if page_index == 0 else "Second page content"
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _build_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Chapter One", level=1)
    document.add_paragraph("Some body text here.")
    document.add_paragraph("A bullet point", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "A2"
    table.cell(1, 0).text = "B1"
    table.cell(1, 1).text = "B2"
    document.add_heading("Chapter Two", level=1)
    document.add_paragraph("More text")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Slide One"
    body = slide.placeholders[1].text_frame
    body.text = "First paragraph"
    paragraph = body.add_paragraph()
    paragraph.text = "Bullet item"
    paragraph.level = 1
    slide.notes_slide.notes_text_frame.text = "Presenter note"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


class TestParseDocument:
    def test_unsupported_extension_raises(self) -> None:
        with pytest.raises(ExtractionError, match="Unsupported document type"):
            parse_document(b"data", "notes.odt")

    def test_dispatches_by_extension(self) -> None:
        pdf_units = parse_document(_build_pdf("Hello World"), "slides.pdf")
        assert pdf_units[0].unit_type == ContentUnitType.SLIDE.value

        txt_units = parse_document(b"Hello\n\nWorld", "notes.txt")
        assert txt_units[0].unit_type == ContentUnitType.DOCUMENT.value


class TestParsePdf:
    def test_single_page(self) -> None:
        units = parse_pdf(_build_pdf("Hello World"), "slides.pdf")
        assert len(units) == 1
        unit = units[0]
        assert unit.unit_type == ContentUnitType.SLIDE.value
        assert unit.position == 1
        assert unit.source_page == 1
        assert unit.title == "Hello World"
        assert unit.raw_text == "Hello World"
        assert len(unit.blocks) == 1
        assert unit.blocks[0].block_type == ContentBlockType.PARAGRAPH.value

    def test_multiple_pages(self) -> None:
        units = parse_pdf(_build_pdf("First slide", pages=2), "slides.pdf")
        assert len(units) == 2
        assert units[0].source_page == 1
        assert units[1].source_page == 2
        assert units[1].raw_text == "Second page content"

    def test_corrupted_file_raises(self) -> None:
        with pytest.raises(ExtractionError, match="Failed to extract"):
            parse_pdf(b"not-a-real-pdf", "slides.pdf")


class TestParseDocx:
    def test_parses_sections_and_blocks(self) -> None:
        units = parse_docx(_build_docx(), "notes.docx")
        assert len(units) == 2
        first = units[0]
        assert first.unit_type == ContentUnitType.SECTION.value
        assert first.title == "Chapter One"
        block_types = [block.block_type for block in first.blocks]
        assert block_types == [
            ContentBlockType.HEADING.value,
            ContentBlockType.PARAGRAPH.value,
            ContentBlockType.LIST_ITEM.value,
            ContentBlockType.TABLE.value,
        ]
        table_block = first.blocks[3]
        assert table_block.metadata == {"rows": 2, "columns": 2}
        assert "A1 | A2" in table_block.content

        second = units[1]
        assert second.title == "Chapter Two"
        assert second.blocks[0].block_type == ContentBlockType.HEADING.value

    def test_corrupted_file_raises(self) -> None:
        with pytest.raises(ExtractionError, match="Failed to extract"):
            parse_docx(b"not-a-docx", "notes.docx")


class TestParsePptx:
    def test_parses_slide_structure(self) -> None:
        units = parse_pptx(_build_pptx(), "deck.pptx")
        assert len(units) == 1
        unit = units[0]
        assert unit.unit_type == ContentUnitType.SLIDE.value
        assert unit.title == "Slide One"
        block_types = [block.block_type for block in unit.blocks]
        assert block_types == [
            ContentBlockType.PARAGRAPH.value,
            ContentBlockType.LIST_ITEM.value,
            ContentBlockType.NOTE.value,
        ]
        assert "First paragraph" in unit.raw_text

    def test_corrupted_file_raises(self) -> None:
        with pytest.raises(ExtractionError, match="Failed to extract"):
            parse_pptx(b"not-a-pptx", "deck.pptx")


class TestParseTxt:
    def test_parses_paragraphs(self) -> None:
        units = parse_txt(b"First paragraph\n\nSecond paragraph", "notes.txt")
        assert len(units) == 1
        unit = units[0]
        assert unit.unit_type == ContentUnitType.DOCUMENT.value
        assert unit.position == 1
        assert len(unit.blocks) == 2
        assert all(block.block_type == ContentBlockType.PARAGRAPH.value for block in unit.blocks)

    def test_empty_content_returns_no_units(self) -> None:
        assert parse_txt(b"", "empty.txt") == []
        assert parse_txt(b"   \n\n  ", "blank.txt") == []

    def test_heading_like_first_line_is_title(self) -> None:
        units = parse_txt(b"Introduction\n\nBody text here", "notes.txt")
        unit = units[0]
        assert unit.title == "Introduction"
        assert unit.blocks[0].block_type == ContentBlockType.HEADING.value
        assert unit.blocks[1].block_type == ContentBlockType.PARAGRAPH.value
