"""Checkpoint C1: Source-Content Integration & PPT Learning Foundation Tests.

Verifies:
1. PPTX document parser retains slide order, titles, paragraphs, hierarchical bullet levels,
   table data with row/column counts, and speaker notes.
2. Full content extraction pipeline persists content_units and content_blocks into PostgreSQL.
3. Authenticated Content API: GET /api/v1/presentations/{id}/content returns complete source content.
4. Player Start API: POST /api/v1/lessons/{id}/player/start returns source_units and presentation metadata.
5. Cross-user isolation: User B cannot access User A's presentation content or player session.
6. Lesson position tracking and slide-accurate resume operate reliably.
"""

from __future__ import annotations

import io
import uuid
from typing import Any
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
import pytest_asyncio
from fastapi import HTTPException, Request
from fastapi import status as http_status
from httpx import AsyncClient
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import select

from app.core.security import create_access_token, hash_password
from app.database.unit_of_work import UnitOfWork
from app.models.content_block import ContentBlock
from app.models.content_unit import ContentUnit
from app.models.generated_lesson import GeneratedLesson
from app.models.generated_lesson_version import GeneratedLessonVersion
from app.models.presentation import Presentation as PresentationModel
from app.models.user import User
from app.parsers.document_parser import parse_pptx
from app.services.content_extraction_service import ContentExtractionService
from app.services.lesson_player_service import LessonPlayerService
from shared.constants import (
    ContentBlockType,
    ContentUnitType,
    ExtractionStatus,
    LessonDifficulty,
    LessonStatus,
    LessonVersionStatus,
)
from tests.conftest import TestSessionLocal

pytestmark = pytest.mark.asyncio

_USER_A_ID = uuid.UUID("a1a1a1a1-a1a1-a1a1-a1a1-a1a1a1a1a1a1")
_USER_B_ID = uuid.UUID("b2b2b2b2-b2b2-b2b2-b2b2-b2b2b2b2b2b2")


def create_educational_pptx() -> bytes:
    """Generate an in-memory PPTX with hierarchical bullets, a structured table, and notes."""
    prs = Presentation()

    # Slide 1: Hierarchical Bullets
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "Computer Networks Overview"
    tf1 = slide1.placeholders[1].text_frame
    tf1.text = "A computer network connects multiple devices to share data."

    p1 = tf1.add_paragraph()
    p1.text = "Local Area Network (LAN)"
    p1.level = 0

    p2 = tf1.add_paragraph()
    p2.text = "Covers small geographic area like school or office"
    p2.level = 1

    p3 = tf1.add_paragraph()
    p3.text = "High speed data transfer rate"
    p3.level = 1

    p4 = tf1.add_paragraph()
    p4.text = "Wide Area Network (WAN)"
    p4.level = 0

    p5 = tf1.add_paragraph()
    p5.text = "Spans large geographical distances across cities"
    p5.level = 1

    # Slide 2: Structured Table & Speaker Notes
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "LAN vs WAN Comparison"

    rows, cols = 3, 3
    left, top, width, height = Inches(1), Inches(2), Inches(8), Inches(2.5)
    table_shape = slide2.shapes.add_table(rows, cols, left, top, width, height)
    tbl = table_shape.table

    table_matrix = [
        ["Metric", "LAN", "WAN"],
        ["Geographic Scope", "Building or campus", "Global or country-wide"],
        ["Transfer Speed", "1000 Mbps+", "Up to 100 Mbps"],
    ]
    for r_idx, row in enumerate(table_matrix):
        for c_idx, val in enumerate(row):
            tbl.cell(r_idx, c_idx).text = val

    notes_slide = slide2.notes_slide
    notes_slide.notes_text_frame.text = "Discuss fiber optics vs satellite links."

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class _RealUser:
    def __init__(self, user_id: uuid.UUID, email: str, name: str) -> None:
        self.id = user_id
        self.email = email
        self.name = name


@pytest_asyncio.fixture(autouse=True)
async def seed_users():
    from fastapi import Request

    from app.core.dependencies import get_current_user
    from app.core.security import decode_access_token
    from app.main import app

    async with TestSessionLocal() as session:
        for uid, email, name in (
            (_USER_A_ID, "learner_a@eduvision.test", "Learner A"),
            (_USER_B_ID, "learner_b@eduvision.test", "Learner B"),
        ):
            res = await session.execute(select(User).where(User.id == uid))
            if res.scalar_one_or_none() is None:
                session.add(
                    User(
                        id=uid,
                        email=email,
                        name=name,
                        password_hash=hash_password("testpass123"),
                    )
                )
        await session.commit()

    app.dependency_overrides.pop(get_current_user, None)

    user_map = {
        str(_USER_A_ID): _RealUser(_USER_A_ID, "learner_a@eduvision.test", "Learner A"),
        str(_USER_B_ID): _RealUser(_USER_B_ID, "learner_b@eduvision.test", "Learner B"),
    }

    async def _jwt_get_current_user(request: Request) -> Any:
        token = None
        auth_header = request.headers.get("authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
        if not token:
            token = request.cookies.get("access_token")
        if not token:
            raise HTTPException(
                status_code=http_status.HTTP_401_UNAUTHORIZED,
                detail="Not authenticated.",
            )
        try:
            payload = decode_access_token(token)
        except Exception:
            raise HTTPException(
                status_code=http_status.HTTP_401_UNAUTHORIZED,
                detail="Invalid or expired token.",
            )
        user = user_map.get(payload.get("sub"))
        if not user:
            raise HTTPException(
                status_code=http_status.HTTP_401_UNAUTHORIZED,
                detail="User not found.",
            )
        return user

    app.dependency_overrides[get_current_user] = _jwt_get_current_user
    yield
    app.dependency_overrides.pop(get_current_user, None)


@pytest.fixture
def auth_headers_user_a() -> dict[str, str]:
    token = create_access_token(_USER_A_ID, role="user")
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def auth_headers_user_b() -> dict[str, str]:
    token = create_access_token(_USER_B_ID, role="user")
    return {"Authorization": f"Bearer {token}"}


# ── Test Suite ────────────────────────────────────────────────────────────

class TestC1SourceIntegration:
    async def test_pptx_parser_fidelity(self):
        """Verify PPTX parser preserves titles, paragraphs, bullet levels, and table data."""
        pptx_bytes = create_educational_pptx()
        units = parse_pptx(pptx_bytes, "networks.pptx")

        assert len(units) == 2, "Must extract exactly 2 slides"

        # Slide 1 Verification
        s1 = units[0]
        assert s1.title == "Computer Networks Overview"
        assert s1.position == 1
        assert s1.unit_type == ContentUnitType.SLIDE.value

        bullet_blocks = [b for b in s1.blocks if b.block_type == ContentBlockType.LIST_ITEM.value]
        assert len(bullet_blocks) >= 4, "Must preserve bullet items"

        # Hierarchy verification: level 0 and level 1
        levels = [b.metadata.get("level") for b in bullet_blocks if b.metadata]
        assert 0 in levels, "Level 0 bullets must be preserved"
        assert 1 in levels, "Level 1 nested bullets must be preserved"

        # Slide 2 Verification (Table & Notes)
        s2 = units[1]
        assert s2.title == "LAN vs WAN Comparison"
        assert s2.position == 2

        table_blocks = [b for b in s2.blocks if b.block_type == ContentBlockType.TABLE.value]
        assert len(table_blocks) == 1, "Must extract table block"
        tbl_block = table_blocks[0]
        assert tbl_block.metadata is not None
        assert tbl_block.metadata["rows"] == 3
        assert tbl_block.metadata["columns"] == 3
        assert len(tbl_block.metadata["table_data"]) == 3
        assert tbl_block.metadata["table_data"][0] == ["Metric", "LAN", "WAN"]

        note_blocks = [b for b in s2.blocks if b.block_type == ContentBlockType.NOTE.value]
        assert len(note_blocks) == 1, "Must extract speaker notes"
        assert "fiber optics" in note_blocks[0].content.lower()

    async def test_full_source_pipeline_and_player_state(
        self,
        client: AsyncClient,
        auth_headers_user_a: dict[str, str],
        auth_headers_user_b: dict[str, str],
    ):
        """Verify end-to-end extraction persistence, player start, and user isolation."""
        pptx_bytes = create_educational_pptx()

        # 1. User A creates a presentation and uploads source
        pres_res = await client.post(
            "/api/v1/presentations",
            json={"title": "Computer Networks Deep Dive"},
            headers=auth_headers_user_a,
        )
        assert pres_res.status_code == 201, pres_res.text
        pres_id = pres_res.json()["data"]["id"]

        # Mock storage for upload and download
        mock_storage = MagicMock()
        mock_storage.upload_fileobj = AsyncMock(return_value="sources/networks.pptx")
        mock_storage.download_fileobj = AsyncMock(return_value=pptx_bytes)

        with patch(
            "app.services.presentation_service.get_storage_backend",
            AsyncMock(return_value=mock_storage),
        ), patch(
            "app.services.content_extraction_service.get_storage_backend",
            AsyncMock(return_value=mock_storage),
        ):
            up_res = await client.post(
                f"/api/v1/presentations/{pres_id}/source",
                files={"source": ("networks.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
                headers=auth_headers_user_a,
            )
            assert up_res.status_code == 200

            # Trigger extraction service
            async with UnitOfWork() as uow:
                await ContentExtractionService(uow).extract_presentation(pres_id)

        # 2. Verify GET /api/v1/presentations/{id}/content for User A
        content_res = await client.get(
            f"/api/v1/presentations/{pres_id}/content",
            headers=auth_headers_user_a,
        )
        assert content_res.status_code == 200
        content_data = content_res.json()["data"]
        assert len(content_data["units"]) == 2
        units = content_data["units"]
        assert units[0]["title"] == "Computer Networks Overview"
        assert units[1]["title"] == "LAN vs WAN Comparison"

        # Check table block survives through content API
        table_blocks = [b for b in units[1]["blocks"] if b["block_type"] == "table"]
        assert len(table_blocks) == 1
        assert table_blocks[0]["metadata"]["table_data"][0] == ["Metric", "LAN", "WAN"]

        # 3. Create a generated lesson linked to presentation
        async with UnitOfWork() as uow:
            from app.repositories.presentation_repository import PresentationRepository
            pres_repo = PresentationRepository(uow.session)
            pres_obj = await pres_repo.get_by_public_id_or_raise(pres_id)

            lesson = GeneratedLesson(
                user_id=_USER_A_ID,
                presentation_id=pres_obj.id,
                title="Computer Networks Lesson",
                mode="lesson",
                language="en",
                difficulty=LessonDifficulty.INTERMEDIATE.value,
                status=LessonStatus.READY.value,
                latest_version=1,
            )
            uow.session.add(lesson)
            await uow.flush()

            version = GeneratedLessonVersion(
                lesson_id=lesson.id,
                version=1,
                status=LessonVersionStatus.SUCCEEDED.value,
                title=lesson.title,
                summary="AI Generated Summary",
                language="en",
                difficulty=LessonDifficulty.INTERMEDIATE.value,
                model="test-model",
            )
            uow.session.add(version)
            await uow.flush()
            await uow.commit()
            lesson_public_id = lesson.public_id

        # 4. Verify POST /api/v1/lessons/{id}/player/start returns source_units
        start_res = await client.post(
            f"/api/v1/lessons/{lesson_public_id}/player/start",
            json={},
            headers=auth_headers_user_a,
        )
        assert start_res.status_code == 200
        player_state = start_res.json()["data"]

        assert "source_units" in player_state, "player state must include source_units"
        assert len(player_state["source_units"]) == 2
        assert player_state["source_units"][0]["title"] == "Computer Networks Overview"
        assert player_state["source_units"][1]["title"] == "LAN vs WAN Comparison"

        # Verify presentation info included
        assert "presentation" in player_state
        assert player_state["presentation"]["id"] == pres_id
        assert player_state["presentation"]["slide_count"] == 2

        # 5. Verify Cross-User Isolation (User B cannot access User A's presentation or player)
        cross_content = await client.get(
            f"/api/v1/presentations/{pres_id}/content",
            headers=auth_headers_user_b,
        )
        assert cross_content.status_code in (403, 404), "User B must NOT access User A presentation"

        cross_player = await client.post(
            f"/api/v1/lessons/{lesson_public_id}/player/start",
            json={},
            headers=auth_headers_user_b,
        )
        assert cross_player.status_code in (403, 404), "User B must NOT start User A lesson player"

        # 6. Verify Position Sync & Resume for User A
        session_id = player_state["session"]["session_id"]
        pos_res = await client.post(
            f"/api/v1/lessons/{lesson_public_id}/player/position",
            json={"session_id": session_id, "slide_index": 1},
            headers=auth_headers_user_a,
        )
        assert pos_res.status_code == 200

        # Subsequent GET player resumes slide 1
        state_res = await client.get(
            f"/api/v1/lessons/{lesson_public_id}/player",
            headers=auth_headers_user_a,
        )
        assert state_res.status_code == 200
        assert state_res.json()["data"]["session"]["slide_index"] == 1
