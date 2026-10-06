"""Script to generate a rich educational PPTX for Checkpoint C1 verification."""
import os

from pptx import Presentation
from pptx.util import Inches


def build_presentation(out_path: str) -> None:
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    prs = Presentation()

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Computer Networks & Architecture"
    if len(s1.placeholders) > 1:
        s1.placeholders[1].text = "Unit 1: Fundamentals of Network Communications"

    # Slide 2: Hierarchical Bullets
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Network Types & Classifications"
    tf2 = s2.placeholders[1].text_frame
    tf2.text = "Networks are categorized based on geographic scale and infrastructure:"

    p = tf2.add_paragraph()
    p.text = "Local Area Network (LAN)"
    p.level = 0

    p = tf2.add_paragraph()
    p.text = "Covers a small area such as a home, office, or university campus"
    p.level = 1

    p = tf2.add_paragraph()
    p.text = "High data transfer rates (1 Gbps to 10 Gbps) with low latency"
    p.level = 1

    p = tf2.add_paragraph()
    p.text = "Metropolitan Area Network (MAN)"
    p.level = 0

    p = tf2.add_paragraph()
    p.text = "Spans a city or regional district connecting enterprise branches"
    p.level = 1

    p = tf2.add_paragraph()
    p.text = "Wide Area Network (WAN)"
    p.level = 0

    p = tf2.add_paragraph()
    p.text = "Spans national or global distances using telco backbones and fiber"
    p.level = 1

    # Slide 3: OSI vs TCP/IP Comparison Table
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "OSI vs TCP/IP Protocol Architecture"

    rows, cols = 4, 3
    left, top, width, height = Inches(0.8), Inches(1.8), Inches(8.4), Inches(3.2)
    tbl_shape = s3.shapes.add_table(rows, cols, left, top, width, height)
    tbl = tbl_shape.table

    matrix = [
        ["Layer Tier", "OSI 7-Layer Reference", "TCP/IP 4-Layer Suite"],
        ["Application & UI", "Application, Presentation, Session", "Application (HTTP, DNS, SSH)"],
        ["End-to-End Transport", "Transport Layer (TCP, UDP flow control)", "Transport Layer (TCP, UDP)"],
        ["Packet Routing", "Network Layer (Logical IP addressing)", "Internet Layer (IPv4, IPv6, ICMP)"],
    ]

    for r_idx, row in enumerate(matrix):
        for c_idx, cell_value in enumerate(row):
            tbl.cell(r_idx, c_idx).text = cell_value

    notes3 = s3.notes_slide
    notes3.notes_text_frame.text = "Instructor Note: Clarify that OSI is a conceptual model, whereas TCP/IP is the practical Internet protocol suite."

    # Slide 4: Transmission Media Comparison Table
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Physical Transmission Media Comparison"

    rows, cols = 4, 4
    left, top, width, height = Inches(0.6), Inches(1.8), Inches(8.8), Inches(3.2)
    tbl_shape4 = s4.shapes.add_table(rows, cols, left, top, width, height)
    tbl4 = tbl_shape4.table

    matrix4 = [
        ["Media Category", "Typical Bandwidth", "Max Distance", "Noise Immunity"],
        ["Cat6a Twisted Pair", "10 Gbps", "100 meters", "Moderate (UTP/STP)"],
        ["Coaxial Cable", "1 Gbps", "500 meters", "Good shielding"],
        ["Single-mode Fiber", "100+ Gbps", "40+ kilometers", "Immune (Light pulses)"],
    ]

    for r_idx, row in enumerate(matrix4):
        for c_idx, cell_value in enumerate(row):
            tbl4.cell(r_idx, c_idx).text = cell_value

    prs.save(out_path)
    print(f"Generated sample presentation at {out_path}")

if __name__ == "__main__":
    target = os.path.join(os.path.dirname(__file__), "..", "tests", "fixtures", "computer_networks_sample.pptx")
    build_presentation(os.path.abspath(target))
