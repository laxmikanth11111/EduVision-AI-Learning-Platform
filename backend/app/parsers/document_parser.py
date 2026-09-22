"""Structured text extraction from presentation source documents.

Parsers convert uploaded source files (PDF, DOCX, PPTX, TXT) into a list of
:class:`ExtractedUnit` objects. Each unit is a document section (slide, page,
section or the whole document) and contains an ordered list of typed
:class:`ExtractedBlock` records.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from app.core.exceptions import ExtractionError
from app.utils.file_helpers import get_file_extension
from shared.constants import ContentBlockType, ContentUnitType

_HEADING_STYLE_NAMES = {
    "title",
    "subtitle",
    "heading 1",
    "heading 2",
    "heading 3",
    "heading 4",
    "heading 5",
    "heading 6",
}
_HEADING_LEVELS = {
    "heading 1": 1,
    "heading 2": 2,
    "heading 3": 3,
    "heading 4": 4,
    "heading 5": 5,
    "heading 6": 6,
}
_TERMINAL_PUNCTUATION = {".", "!", "?", ":", ";", ","}


@dataclass
class ExtractedBlock:
    block_type: str
    content: str | None = None
    metadata: dict[str, Any] | None = None


@dataclass
class ExtractedUnit:
    unit_type: str
    position: int
    title: str | None = None
    raw_text: str | None = None
    source_page: int | None = None
    blocks: list[ExtractedBlock] = field(default_factory=list)


def parse_document(content: bytes, filename: str) -> list[ExtractedUnit]:
    """Parse an uploaded source document into structured content units."""
    extension = get_file_extension(filename).lower()
    if extension == ".pdf":
        return parse_pdf(content, filename)
    if extension == ".docx":
        return parse_docx(content, filename)
    if extension == ".pptx":
        return parse_pptx(content, filename)
    if extension == ".txt":
        return parse_txt(content, filename)
    raise ExtractionError(
        message="Unsupported document type for extraction",
        details={"extension": extension or None, "filename": filename},
    )


def parse_pdf(content: bytes, filename: str) -> list[ExtractedUnit]:
    try:
        from pypdf import PasswordType, PdfReader

        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            try:
                if reader.decrypt("") == PasswordType.NOT_DECRYPTED:
                    raise ExtractionError(
                        message=(
                            "This PDF is password-protected. "
                            "Please remove password protection and re-upload."
                        )
                    )
            except ExtractionError:
                raise
            except Exception as exc:
                raise ExtractionError(
                    message=(
                        "This PDF is password-protected. "
                        "Please remove password protection and re-upload."
                    )
                ) from exc

        units: list[ExtractedUnit] = []
        for page_number, page in enumerate(reader.pages, start=1):
            raw_text = (page.extract_text() or "").strip()
            blocks = _text_to_paragraph_blocks(raw_text)
            title = blocks[0].content if blocks else None
            units.append(
                ExtractedUnit(
                    unit_type=ContentUnitType.SLIDE.value,
                    position=page_number,
                    title=title,
                    raw_text=raw_text or None,
                    source_page=page_number,
                    blocks=blocks,
                )
            )
        return units
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            message="Failed to extract content from the PDF file",
            details={"filename": filename, "error": str(exc)[:500]},
        ) from exc


def parse_docx(content: bytes, filename: str) -> list[ExtractedUnit]:
    try:
        from docx import Document
        from docx.table import Table

        document = Document(io.BytesIO(content))
        units: list[ExtractedUnit] = []
        current_blocks: list[ExtractedBlock] = []
        current_title: str | None = None

        def _flush() -> None:
            nonlocal current_blocks, current_title
            if current_blocks:
                raw_text = "\n".join(block.content for block in current_blocks if block.content)
                units.append(
                    ExtractedUnit(
                        unit_type=ContentUnitType.SECTION.value,
                        position=len(units) + 1,
                        title=current_title,
                        raw_text=raw_text or None,
                        blocks=current_blocks,
                    )
                )
            current_blocks = []
            current_title = None

        for item in document.iter_inner_content():
            if isinstance(item, Table):
                rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
                block_content = "\n".join(" | ".join(row) for row in rows)
                if block_content:
                    current_blocks.append(
                        ExtractedBlock(
                            block_type=ContentBlockType.TABLE.value,
                            content=block_content,
                            metadata={
                                "rows": len(rows),
                                "columns": len(rows[0]) if rows else 0,
                            },
                        )
                    )
                continue

            text = (item.text or "").strip()
            if not text:
                continue
            style_name = (item.style.name if item.style else "").lower()
            if style_name in _HEADING_STYLE_NAMES:
                level = _HEADING_LEVELS.get(style_name, 1)
                if level <= 1 or (style_name == "title" and not current_blocks):
                    _flush()
                if current_title is None:
                    current_title = text
                current_blocks.append(
                    ExtractedBlock(
                        block_type=ContentBlockType.HEADING.value,
                        content=text,
                        metadata={"level": level},
                    )
                )
            elif style_name.startswith("list"):
                current_blocks.append(
                    ExtractedBlock(
                        block_type=ContentBlockType.LIST_ITEM.value,
                        content=text,
                    )
                )
            else:
                current_blocks.append(
                    ExtractedBlock(
                        block_type=ContentBlockType.PARAGRAPH.value,
                        content=text,
                    )
                )
        _flush()
        return units
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            message="Failed to extract content from the DOCX file",
            details={"filename": filename, "error": str(exc)[:500]},
        ) from exc


def parse_pptx(content: bytes, filename: str) -> list[ExtractedUnit]:
    try:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(content))
        units: list[ExtractedUnit] = []
        for index, slide in enumerate(prs.slides, start=1):
            title: str | None = None
            if slide.shapes.title is not None:
                title_text = slide.shapes.title.text.strip()
                title = title_text or None

            blocks: list[ExtractedBlock] = []
            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if getattr(shape, "has_text_frame", False) and shape.text_frame:
                    for para in shape.text_frame.paragraphs:
                        para_text = ("".join(run.text for run in para.runs) or getattr(para, "text", "") or "").strip()
                        if not para_text:
                            continue
                        level = getattr(para, "level", 0) or 0
                        is_bullet = (
                            level > 0
                            or (getattr(shape, "is_placeholder", False) and getattr(shape, "placeholder_format", None) is not None and getattr(shape.placeholder_format, "idx", 0) != 0)
                            or any(para_text.startswith(bullet_prefix) for bullet_prefix in ("•", "-", "*", "–", "—", "▪", "▫", "►"))
                        )
                        if is_bullet:
                            blocks.append(
                                ExtractedBlock(
                                    block_type=ContentBlockType.LIST_ITEM.value,
                                    content=para_text,
                                    metadata={"level": level},
                                )
                            )
                        else:
                            blocks.append(
                                ExtractedBlock(
                                    block_type=ContentBlockType.PARAGRAPH.value,
                                    content=para_text,
                                    metadata={"level": level},
                                )
                            )
                elif getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                    block_content = "\n".join(" | ".join(row) for row in rows)
                    blocks.append(
                        ExtractedBlock(
                            block_type=ContentBlockType.TABLE.value,
                            content=block_content,
                            metadata={
                                "rows": len(rows),
                                "columns": len(rows[0]) if rows else 0,
                                "table_data": rows,
                            },
                        )
                    )
                elif getattr(shape, "shape_type", None) is not None and str(shape.shape_type) in {
                    "PICTURE (13)",
                    "13",
                }:
                    blocks.append(
                        ExtractedBlock(
                            block_type=ContentBlockType.IMAGE.value,
                            metadata={"shape_name": getattr(shape, "name", "")},
                        )
                    )

            notes_text: str | None = None
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text.strip() or None
            if notes_text:
                blocks.append(
                    ExtractedBlock(
                        block_type=ContentBlockType.NOTE.value,
                        content=notes_text,
                    )
                )

            text_blocks = [b for b in blocks if b.content]
            contents = [str(b.content) for b in text_blocks]
            raw_text = "\n".join([title or "", *contents]).strip()
            units.append(
                ExtractedUnit(
                    unit_type=ContentUnitType.SLIDE.value,
                    position=index,
                    title=title,
                    raw_text=raw_text or None,
                    source_page=index,
                    blocks=blocks,
                )
            )
        return units
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(
            message="Failed to extract content from the PPTX file",
            details={"filename": filename, "error": str(exc)[:500]},
        ) from exc


def parse_txt(content: bytes, filename: str) -> list[ExtractedUnit]:
    try:
        text = content.decode("utf-8", errors="replace")
        lines = [line.strip() for line in text.splitlines()]
        paragraphs: list[str] = []
        buffer: list[str] = []
        for line in lines:
            if not line:
                if buffer:
                    paragraphs.append("\n".join(buffer))
                    buffer = []
                continue
            buffer.append(line)
        if buffer:
            paragraphs.append("\n".join(buffer))

        if not paragraphs:
            return []

        blocks: list[ExtractedBlock] = []
        title: str | None = None
        for index, paragraph in enumerate(paragraphs):
            if _is_heading_like(paragraph) and len(paragraph.splitlines()) == 1:
                if index == 0 and title is None:
                    title = paragraph
                blocks.append(
                    ExtractedBlock(
                        block_type=ContentBlockType.HEADING.value,
                        content=paragraph,
                    )
                )
            else:
                blocks.append(
                    ExtractedBlock(
                        block_type=ContentBlockType.PARAGRAPH.value,
                        content=paragraph,
                    )
                )

        raw_text = "\n".join(paragraphs)
        return [
            ExtractedUnit(
                unit_type=ContentUnitType.DOCUMENT.value,
                position=1,
                title=title,
                raw_text=raw_text or None,
                blocks=blocks,
            )
        ]
    except Exception as exc:
        raise ExtractionError(
            message="Failed to extract content from the TXT file",
            details={"filename": filename, "error": str(exc)[:500]},
        ) from exc


def _text_to_paragraph_blocks(text: str) -> list[ExtractedBlock]:
    """Split raw text into paragraph blocks on blank lines."""
    blocks: list[ExtractedBlock] = []
    buffer: list[str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            if buffer:
                blocks.append(
                    ExtractedBlock(
                        block_type=ContentBlockType.PARAGRAPH.value,
                        content="\n".join(buffer),
                    )
                )
                buffer = []
            continue
        buffer.append(stripped)
    if buffer:
        blocks.append(
            ExtractedBlock(
                block_type=ContentBlockType.PARAGRAPH.value,
                content="\n".join(buffer),
            )
        )
    return blocks


def _is_heading_like(text: str) -> bool:
    if not text or len(text) > 60 or len(text.splitlines()) != 1:
        return False
    if text[-1] in _TERMINAL_PUNCTUATION:
        return False
    words = text.split()
    if not words:
        return False
    if len(words) == 1:
        return text[0].isupper() or text.isupper()
    return all(word[0].isupper() for word in words)
