from __future__ import annotations

import io

from app.models.export_template import ExportTemplate
from app.schemas.export import ExportContentData, ExportOptions
from app.services.exporters.base_exporter import BaseExporter, ExportResult


class DOCXExporter(BaseExporter):
    async def export(
        self,
        data: ExportContentData,
        options: ExportOptions,
        template: ExportTemplate | None = None,
    ) -> ExportResult:
        try:
            import docx

            doc = docx.Document()
            doc.add_heading(data.title, level=0)

            if data.subtitle:
                p = doc.add_paragraph()
                p.add_run(data.subtitle).italic = True

            meta_parts: list[str] = []
            if data.author:
                meta_parts.append(f"Author: {data.author}")
            if data.created_at:
                meta_parts.append(f"Date: {data.created_at}")
            if meta_parts:
                doc.add_paragraph(" | ".join(meta_parts))

            for section in data.sections:
                sec_title = section.get("title", "Section")
                doc.add_heading(sec_title, level=1)

                content = section.get("content")
                if content:
                    doc.add_paragraph(content)

                for block in section.get("blocks", []):
                    b_heading = block.get("heading")
                    if b_heading:
                        doc.add_heading(b_heading, level=2)
                    b_content = block.get("content", "")
                    if b_content:
                        doc.add_paragraph(str(b_content))

            if data.quiz_questions:
                doc.add_heading("Practice Quiz", level=1)
                for idx, q in enumerate(data.quiz_questions, start=1):
                    doc.add_paragraph(f"Q{idx}. {q.get('question', '')}", style="List Bullet")
                    for opt in q.get("options", []):
                        doc.add_paragraph(f"  • {opt}")
                    if options.include_answer_key and q.get("answer"):
                        p = doc.add_paragraph(f"Answer: {q.get('answer')}")
                        p.runs[0].italic = True

            buffer = io.BytesIO()
            doc.save(buffer)
            docx_bytes = buffer.getvalue()

        except ImportError:
            # Fallback text-based document
            lines = [f"# {data.title}\n"]
            for s in data.sections:
                lines.append(f"## {s.get('title', 'Section')}\n{s.get('content', '')}\n")
            docx_bytes = "\n".join(lines).encode("utf-8")

        safe_title = "".join(c for c in data.title.lower() if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
        filename = f"{safe_title or 'export'}.docx"

        return ExportResult.create(
            content=docx_bytes,
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
