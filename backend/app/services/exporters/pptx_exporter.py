from __future__ import annotations

import io

from app.models.export_template import ExportTemplate
from app.schemas.export import ExportContentData, ExportOptions
from app.services.exporters.base_exporter import BaseExporter, ExportResult


class PPTXExporter(BaseExporter):
    async def export(
        self,
        data: ExportContentData,
        options: ExportOptions,
        template: ExportTemplate | None = None,
    ) -> ExportResult:
        try:
            import pptx

            prs = pptx.Presentation()

            # Title Slide
            title_slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(title_slide_layout)
            title = slide.shapes.title
            subtitle = slide.placeholders[1]

            title.text = data.title
            if data.subtitle:
                subtitle.text = data.subtitle
            elif data.author:
                subtitle.text = f"By {data.author}"

            # Content Slides
            bullet_slide_layout = prs.slide_layouts[1]
            for section in data.sections:
                sec_slide = prs.slides.add_slide(bullet_slide_layout)
                sec_slide.shapes.title.text = section.get("title", "Slide")

                tf = sec_slide.placeholders[1].text_frame
                content = section.get("content")
                if content:
                    tf.text = content

                for block in section.get("blocks", []):
                    b_content = block.get("content", "")
                    if b_content:
                        p = tf.add_paragraph()
                        p.text = str(b_content)
                        p.level = 0

            buffer = io.BytesIO()
            prs.save(buffer)
            pptx_bytes = buffer.getvalue()

        except ImportError:
            # Fallback formatted bytes
            lines = [f"SLIDE 1: {data.title}\n"]
            for idx, s in enumerate(data.sections, start=2):
                lines.append(f"SLIDE {idx}: {s.get('title', 'Slide')}\n{s.get('content', '')}\n")
            pptx_bytes = "\n".join(lines).encode("utf-8")

        safe_title = "".join(c for c in data.title.lower() if c.isalnum() or c in (" ", "_", "-")).strip().replace(" ", "_")
        filename = f"{safe_title or 'export'}.pptx"

        return ExportResult.create(
            content=pptx_bytes,
            filename=filename,
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
